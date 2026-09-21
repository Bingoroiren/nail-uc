# -*- coding: utf-8 -*-
"""
Hệ thống Làm giàu dữ liệu Agency Tuyển dụng & Cho thuê Nhân sự Phần Lan (ĐÃ TỐI ƯU HÓA TỐC ĐỘ & ĐỘ CHÍNH XÁC):
1. Tra cứu Google Maps (hl=fi, headless=False) với bộ lọc tên nghiêm ngặt cho Phần Lan.
2. Fallback sang Bing / DuckDuckGo nếu Maps không có kết quả hoặc không có website (có nhận diện Captcha).
3. Cào Website để tìm Email, SĐT và link Facebook (xử lý triệt để mã %20 và mailto).
4. Nếu Website không có Email, cào tiếp Fanpage Facebook để tìm Email.
5. Tối ưu:
   - Domain Cache & FB Cache: tránh cào lại cùng 1 trang web/fanpage nhiều lần (cho các agency nhiều chi nhánh như Barona, Eezy, SOL...).
   - Nhận diện trang liên hệ chuẩn Phần Lan (/yhteystiedot/, /ota-yhteytta/...) loại trừ các link báo cáo tài chính/informaatio.
   - Ưu tiên cào Web trước nếu đã có sẵn website từ Finder.fi giúp tăng tốc gấp 5 lần.
   - Lưu động 100%, bảo toàn dữ liệu khi bấm Ctrl+C, chống lỗi PermissionError khi mở Excel.
"""

import asyncio
import csv
import json
import os
import re
import sys
import html
import urllib.parse
from bs4 import BeautifulSoup
from curl_cffi.requests import AsyncSession
from playwright.async_api import async_playwright
import openpyxl

# Set Windows Proactor Event Loop Policy
if sys.platform == 'win32':
    try:
        asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())
    except Exception:
        pass

try:
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
except Exception:
    pass

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_CSV = os.path.join(BASE_DIR, "agency phần lan (đã lọc trùng).csv")
OUTPUT_CSV = os.path.join(BASE_DIR, "agency phần lan (đã lọc trùng).csv")
OUTPUT_XLSX = os.path.join(BASE_DIR, "agency phần lan.xlsx")
CACHE_FILE = os.path.join(BASE_DIR, "crawlmail", "cache_maps_enrich_agency_fi.json")

LEGAL_FORMS_FI = [
    'oy', 'ab', 'oyj', 'ky', 'ay', 'tmi', 'osk', 'ry', 'ltd'
]

EMAIL_REGEX = re.compile(r'\b[A-Za-z0-9._%+-]{1,64}@[A-Za-z0-9.-]{1,255}\.[A-Za-z]{2,7}\b')
OBF_EMAIL_REGEX = re.compile(
    r'([A-Za-z0-9._%+-]{2,40})\s*(?:@|\[at\]|\(at\)|\[ät\]|\(ät\)|\s+at\s+|\s+ät\s+)\s*([A-Za-z0-9.-]+\.[A-Za-z]{2,7})',
    re.IGNORECASE
)
INVALID_EXTENSIONS = ('.png', '.jpg', '.jpeg', '.gif', '.svg', '.webp', '.pdf', '.css', '.js', '.ico', '.woff', '.woff2', '.mp4', '.mp3', '.ttf')
IGNORE_DOMAINS = {
    'sentry.io', 'wixpress.com', 'example.com', 'domain.com', 'schema.org', 
    'wordpress.org', 'cloudflare.com', 'google.com', 'facebook.com', 'w3.org', 
    'jsdelivr.net', 'bootstrapcdn.com', 'website.com', 'yourdomain.com', 'fb.com', 'email.fi'
}
DUMMY_EMAILS = {'user@website.com', 'name@domain.com', 'email@domain.com', 'info@domain.com', 'contact@domain.com', 'esimerkki@email.fi', 'etunimi.sukunimi@eezy.fi'}

EXCLUDED_PLATFORMS_FI = {
    # Danh bạ, cổng thông tin doanh nghiệp, tra cứu mã số thuế Phần Lan & Quốc tế
    'finder.fi', 'profinder.fi', 'kauppalehti.fi', 'asiakastieto.fi', 'yritystele.fi', 'proff.fi',
    'fonecta.fi', 'fonecta.com', 'suomi.fi', 'prh.fi', 'ytj.fi', 'vero.fi',
    'yritysopas.fi', 'yrityshaku.fi', 'yritykset.fi', 'yritysfakta.fi', 'taloustutka.fi',
    'almamedia.fi', 'directa.fi', 'eniro.fi', 'eniro.se', '0100100.fi', 'sinunyritys.fi',
    'suomenyritykset.fi', 'suomenyrityshaku.fi', 'yritysrekisteri.fi', 'bisnode.fi',
    'bisnode.com', 'dnb.com', 'kompass.com', 'europages.com', 'infobel.com',
    'firmaspraak.fi', 'tietopalvelut.fi', 'avointieto.fi', 'tulli.fi',

    # Tuyển dụng, việc làm, sàn trung gian, rao vặt
    'duunitori.fi', 'oikotie.fi', 'monster.fi', 'jobly.fi', 'tyomarkkinatori.fi',
    'te-palvelut.fi', 'tori.fi', 'indeed.com', 'fi.indeed.com', 'glassdoor.com',
    'glassdoor.fi', 'jooble.org', 'fi.jooble.org', 'stepstone.se', 'stepstone.de',
    'rekrytointi.com', 'uranus.fi', 'workinfinland.com', 'cv-online.com', 'linkedin.com',
    'staffpoint.fi', 'barona.fi', 'eezy.fi', 'bolt.works', 'vmp.fi', 'adecco.fi', 'manpower.fi',
    
    # Mạng xã hội & Video/Audio
    'facebook.com', 'fb.com', 'instagram.com', 'twitter.com', 'x.com', 'youtube.com',
    'tiktok.com', 'pinterest.com', 'wikipedia.org', 'reddit.com', 'vimeo.com',
    
    # Báo chí, truyền thông, diễn đàn Phần Lan
    'yle.fi', 'is.fi', 'hs.fi', 'iltalehti.fi', 'uusisuomi.fi', 'talouselama.fi',
    'tivi.fi', 'tekniikkatalous.fi', 'mtv.fi', 'mtvuutiset.fi', 'suomi24.fi',
    'vauva.fi', 'helsinginuutiset.fi', 'tamperelainen.fi', 'turkulainen.fi',
    'aamulehti.fi', 'kaleva.fi', 'ksml.fi', 'savonsanomat.fi', 'ess.fi',
    'satakunnankansa.fi', 'lapinkansa.fi', 'karjalainen.fi', 'pohjalainen.fi',
    'leadfeeder.com', 'tripadvisor.com', 'trustpilot.com', 'google.com', 'google.fi',
    'bing.com', 'duckduckgo.com', 'yahoo.com', 'msn.com',
    
    # Nền tảng tạo web / blog miễn phí không có tên miền riêng (nếu là trang chủ nền tảng)
    'wix.com', 'wixsite.com', 'wordpress.com', 'wordpress.org', 'weebly.com', 'squarespace.com',
    'shopify.com', 'myshopify.com', 'site123.me', 'jimdosite.com', 'webnode.fi', 'webnode.com',
    'blogspot.com', 'medium.com', 'github.io', 'sites.google.com'
}

PLATFORM_KEYWORDS_FI = {
    'directory', 'yellowpages', 'yrityshaku', 'rekry', 'tyopaikat', 'duunit',
    'katalog', 'listing', 'yritykset', 'portaali', 'portal', 'rekisteri',
    'tietokanta', 'uutiset', 'media', 'sanomat', 'lehti', 'forum', 'keskustelu',
    'arvostelut', 'reviews', 'ratings'
}

GENERIC_NAME_TOKENS = {
    'suomi', 'finland', 'palvelut', 'palvelu', 'group', 'nordic', 'holding', 
    'consulting', 'management', 'international', 'services', 'service', 'team', 
    'staff', 'work', 'works', 'yhtio', 'partner', 'partners', 'henkilosto', 'rekrytointi'
}

# Caches in-memory to prevent re-scraping identical parent domains/Facebook pages
DOMAIN_CACHE = {}
FB_CACHE = {}

def clean_company_name_fi(name):
    if not name:
        return ""
    n = re.sub(r'["\'„“”«»]', '', name).strip().lower()
    tokens = n.split()
    tokens = [t for t in tokens if t not in LEGAL_FORMS_FI]
    return ' '.join(tokens).strip()

def is_valid_name_match_fi(query_name, candidate_name):
    q = clean_company_name_fi(query_name)
    c = re.sub(r'["\'„“”«»]', '', candidate_name).strip().lower()
    c_clean = clean_company_name_fi(candidate_name)
    
    if not q or not c_clean:
        return False, "EMPTY"

    if q == c_clean:
        return True, "EXACT"
        
    c_without_brackets = re.sub(r'\(.*?\)|\[.*?\]', '', c).strip()
    c_without_brackets_clean = clean_company_name_fi(c_without_brackets)
    if q == c_without_brackets_clean:
        return True, "BRACKET_MATCH"
        
    parts = re.split(r'[-–—,.:|/]', c)
    parts_clean = [clean_company_name_fi(p) for p in parts if p.strip()]
    if any(p == q for p in parts_clean):
        return True, "SEPARATOR_MATCH"
        
    return False, "NO_MATCH"

def decode_cloudflare_email(hex_str):
    try:
        hex_data = bytes.fromhex(hex_str)
        key = hex_data[0]
        return ''.join(chr(b ^ key) for b in hex_data[1:])
    except Exception:
        return ""

def clean_email(email_str):
    if not email_str:
        return ""
    em = urllib.parse.unquote(email_str).replace('%20', '').strip().lower().rstrip('.,;:')
    if any(em.endswith(ext) for ext in INVALID_EXTENSIONS):
        return ""
    if em in DUMMY_EMAILS:
        return ""
    domain = em.split('@')[-1]
    if domain in IGNORE_DOMAINS or domain.endswith('.invalid'):
        return ""
    if len(em) < 6 or '@' not in em:
        return ""
    return em

def extract_emails_from_html(html_str):
    if not html_str:
        return set()
    found = set()
    unescaped = html.unescape(html_str)
    
    # 1. Cloudflare emails
    for cf in re.findall(r'data-cfemail="([0-9a-fA-F]+)"', html_str):
        dec = decode_cloudflare_email(cf)
        em = clean_email(dec)
        if em:
            found.add(em)
            
    # 2. mailto: links (xử lý cả %20)
    for mailto in re.findall(r'mailto:([^\s"\'<>]+)', unescaped, re.IGNORECASE):
        cleaned_m = mailto.split('?')[0]
        for sub_em in re.findall(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,7}', urllib.parse.unquote(cleaned_m)):
            em = clean_email(sub_em)
            if em:
                found.add(em)
                
    # 3. Plain text regex
    for match in EMAIL_REGEX.findall(unescaped):
        em = clean_email(match)
        if em:
            found.add(em)
            
    # 4. Obfuscated emails ([at], (at), [ät], (ät), at, ät) - phổ biến ở Phần Lan
    for u_part, d_part in OBF_EMAIL_REGEX.findall(unescaped):
        candidate = f"{u_part.strip()}@{d_part.strip()}"
        em = clean_email(candidate)
        if em:
            found.add(em)
            
    return found

def extract_facebook_url(html):
    if not html:
        return ""
    for match in re.findall(r'https?://(?:www\.)?facebook\.com/(?:[a-zA-Z0-9.\-_]+)', html, re.IGNORECASE):
        m_lower = match.lower()
        if not any(x in m_lower for x in ['/sharer', '/share.php', '/tr', '/dialog', '/plugins', '/hashtag', '/pages', '/profile.php', '/people', '/groups', '/login', '/home.php']):
            m_clean = urllib.parse.unquote(match).split('?')[0].rstrip('/')
            last_slug = m_clean.split('/')[-1].lower()
            if len(last_slug) > 2 and last_slug not in ['people', 'profile.php', 'pages', 'groups', 'login']:
                return m_clean
    return ""

def score_email(email, domain):
    score = 0
    clean_domain = domain.lower().replace('www.', '')
    if clean_domain and clean_domain in email:
        score += 50
    pfx = email.split('@')[0].lower()
    if pfx in ['info', 'rekry', 'rekrytointi', 'asiakaspalvelu', 'toimisto', 'myynti', 'hakemukset', 'contact']:
        score += 30
    elif any(pfx.startswith(x) for x in ['info', 'rekry', 'sales', 'contact', 'ura']):
        score += 20
    return score

def prioritize_emails(emails_set, domain=""):
    if not emails_set:
        return ""
    sorted_emails = sorted(list(emails_set), key=lambda e: score_email(e, domain), reverse=True)
    return "; ".join(sorted_emails)

async def check_for_captcha(page, platform_name="Google"):
    url = page.url.lower()
    content = ""
    try:
        content = await page.content()
        content_lower = content.lower()
    except Exception:
        content_lower = ""
        
    is_blocked = False
    if "sorry/index" in url or "google.com/sorry" in url:
        is_blocked = True
    elif "unusual traffic" in content_lower or "recaptcha" in content_lower and "g-recaptcha" in content_lower:
        is_blocked = True
    elif "detected unusual traffic" in content_lower:
        is_blocked = True

    if is_blocked:
        print("\n" + "!" * 75, flush=True)
        print(f" [CẢNH BÁO BOT] {platform_name} phát hiện hoạt động bất thường / Yêu cầu Captcha!", flush=True)
        print(" Trình duyệt Chrome đang mở sẵn (headless=False) trên màn hình của bạn.", flush=True)
        print(" Vui lòng thao tác giải Captcha trên cửa sổ trình duyệt.", flush=True)
        print(" Sau khi hoàn tất, quay lại đây nhấn phím [ENTER] để tiếp tục cào...", flush=True)
        print("!" * 75 + "\n", flush=True)
        loop = asyncio.get_event_loop()
        await loop.run_in_executor(None, input, ">> Nhấn [ENTER] khi bạn đã giải quyết Captcha: ")
        await asyncio.sleep(2)
        return True
    return False

async def search_google_maps_fi(page, company_name):
    """
    Tra cứu Google Maps với ngôn ngữ Phần Lan (hl=fi).
    Từ khóa tra cứu đúng nguyên văn tên công ty, không thêm bớt.
    Có cơ chế chờ thông minh thích ứng với mạng VPN (chờ kết quả thực tế xuất hiện).
    """
    query = company_name.strip()
    q_clean = clean_company_name_fi(company_name)
    url = f"https://www.google.com/maps/search/{urllib.parse.quote(query)}?hl=fi"
    
    res = {
        "tag_maps": "",
        "phone": "",
        "website": "",
        "address": "",
        "maps_title": "",
        "matched": False,
        "match_reason": ""
    }
    
    try:
        # 1. Điều hướng với timeout dài (25s) thích ứng với mạng VPN
        await page.goto(url, wait_until="domcontentloaded", timeout=25000)
        
        # Xử lý Cookie consent Phần Lan nếu có
        try:
            consent_btn = page.locator('button[aria-label*="Hyväksy"], button[aria-label*="Hylkää"], button[aria-label*="Accept"], form[action*="consent"] button')
            if await consent_btn.count() > 0:
                await consent_btn.first.click()
                await asyncio.sleep(1)
        except Exception:
            pass
            
        await check_for_captcha(page, "Google Maps")
        
        # 2. CHỜ THÔNG MINH CHO VPN: Chờ một trong các dấu hiệu tải xong của Google Maps:
        # - Hoặc thẻ chi tiết doanh nghiệp: h1.DUwDvf, h1[class*="fontHeadlineLarge"]
        # - Hoặc danh sách kết quả tìm kiếm: a.hfpxzc, div[role="feed"]
        # - Hoặc thông báo không tìm thấy kết quả: text="Ei tuloksia", text="tuloksia"
        try:
            await page.wait_for_selector(
                'h1.DUwDvf, h1[class*="fontHeadlineLarge"], a.hfpxzc, div[role="feed"], div.fontBodyMedium, div:has-text("tuloksia")',
                timeout=12000
            )
        except Exception:
            # Nếu mạng VPN chậm, chờ thêm 2.5s
            await asyncio.sleep(2.5)
            
        # Thêm 1.5s để các thành phần SĐT, website, địa chỉ hydrate xong trên DOM
        await asyncio.sleep(1.5)
        
        # Case A: Direct Detail View (Trực tiếp mở trang chi tiết)
        title_elem = page.locator('h1.DUwDvf, h1[class*="fontHeadlineLarge"]')
        if await title_elem.count() > 0:
            title = await title_elem.first.text_content()
            matched, reason = is_valid_name_match_fi(q_clean, title or "")
            if matched:
                res["matched"] = True
                res["match_reason"] = reason
                res["maps_title"] = (title or "").strip()
                
                cat_elem = page.locator('button.DkEaL, button[jsaction*="category"], span.DkEaL')
                if await cat_elem.count() > 0:
                    res["tag_maps"] = (await cat_elem.first.text_content()).strip()
                    
                phone_elem = page.locator('button[data-item-id*="phone"]')
                if await phone_elem.count() > 0:
                    raw_ph = await phone_elem.first.text_content()
                    clean_p = re.sub(r'[\ue000-\uf8ff]', '', raw_ph)
                    res["phone"] = re.sub(r'[^\d+ ]', '', clean_p).strip()
                    
                web_elem = page.locator('a[data-item-id="authority"]')
                if await web_elem.count() > 0:
                    res["website"] = (await web_elem.first.get_attribute("href") or "").strip()
                    
                addr_elem = page.locator('button[data-item-id="address"]')
                if await addr_elem.count() > 0:
                    raw_addr = await addr_elem.first.text_content()
                    clean_a = re.sub(r'[\ue000-\uf8ff]', '', raw_addr)
                    res["address"] = clean_a.strip()
                return res

        # Case B: List of Results (Danh sách nhiều thẻ kết quả)
        cards = page.locator('a.hfpxzc')
        card_count = await cards.count()
        if card_count > 0:
            for i in range(min(card_count, 4)):
                card_title = await cards.nth(i).get_attribute('aria-label')
                matched, reason = is_valid_name_match_fi(q_clean, card_title or "")
                if matched:
                    res["matched"] = True
                    res["match_reason"] = reason
                    res["maps_title"] = (card_title or "").strip()
                    
                    await cards.nth(i).click()
                    # Chờ panel chi tiết mở ra
                    try:
                        await page.wait_for_selector('h1.DUwDvf, button[data-item-id*="phone"], a[data-item-id="authority"]', timeout=6000)
                    except Exception:
                        await asyncio.sleep(2)
                        
                    await asyncio.sleep(1)
                    
                    cat_elem = page.locator('button.DkEaL, button[jsaction*="category"], span.DkEaL')
                    if await cat_elem.count() > 0:
                        res["tag_maps"] = (await cat_elem.first.text_content()).strip()
                        
                    phone_elem = page.locator('button[data-item-id*="phone"]')
                    if await phone_elem.count() > 0:
                        raw_ph = await phone_elem.first.text_content()
                        clean_p = re.sub(r'[\ue000-\uf8ff]', '', raw_ph)
                        res["phone"] = re.sub(r'[^\d+ ]', '', clean_p).strip()
                        
                    web_elem = page.locator('a[data-item-id="authority"]')
                    if await web_elem.count() > 0:
                        res["website"] = (await web_elem.first.get_attribute("href") or "").strip()
                        
                    addr_elem = page.locator('button[data-item-id="address"]')
                    if await addr_elem.count() > 0:
                        raw_addr = await addr_elem.first.text_content()
                        clean_a = re.sub(r'[\ue000-\uf8ff]', '', raw_addr)
                        res["address"] = clean_a.strip()
                    return res
    except Exception:
        pass
        
    return res

def normalize_fi_domain_token(text):
    """Chuyển đổi ký tự tiếng Phần Lan ä->a, ö->o, å->a để so khớp với domain ASCII"""
    t = text.lower()
    t = t.replace('ä', 'a').replace('ö', 'o').replace('å', 'a')
    return re.sub(r'[^a-z0-9]', '', t)

def is_strictly_company_domain_fi(domain, company_name):
    """
    Kiểm tra tên miền có thực sự là WEB RIÊNG CỦA DOANH NGHIỆP hay không:
    - Loại bỏ 100% các trang nền tảng chung, danh bạ, trang tuyển dụng, mạng xã hội, báo chí.
    - Bắt buộc domain phải chứa từ khóa nhận diện đặc thù của thương hiệu công ty.
    """
    if not domain:
        return False
    d = domain.lower().replace('www.', '').strip()
    
    # 1. Trực tiếp nằm trong danh sách đen các nền tảng/danh bạ
    for plat in EXCLUDED_PLATFORMS_FI:
        if d == plat or d.endswith('.' + plat):
            return False
            
    q_clean = clean_company_name_fi(company_name)
    tokens = q_clean.split()
    
    # 2. Kiểm tra từ khóa nền tảng/danh bạ (chỉ cấm nếu từ khóa đó không nằm trong tên cty)
    for kw in PLATFORM_KEYWORDS_FI:
        if kw in d and kw not in q_clean:
            return False
            
    # 3. Phải chứa từ khóa thương hiệu đặc trưng của công ty
    distinctive_tokens = [t for t in tokens if len(t) >= 3 and t not in GENERIC_NAME_TOKENS]
    
    # Lấy phần định danh chính của domain (bỏ TLD)
    d_clean = re.sub(r'[^a-z0-9]', '', d.split('.')[0])
    
    if distinctive_tokens:
        for t in distinctive_tokens:
            t_norm = normalize_fi_domain_token(t)
            if len(t_norm) >= 3 and (t_norm in d or t_norm in d_clean):
                return True
            
    # Nếu toàn từ thông dụng (vd: Nordic Staff Oy -> tokens: nordic, staff)
    clean_nospace = ''.join(normalize_fi_domain_token(t) for t in tokens)
    if len(clean_nospace) >= 5 and clean_nospace in d:
        return True
        
    # Ghép 2 từ đầu
    if len(tokens) >= 2:
        combo = normalize_fi_domain_token(tokens[0]) + normalize_fi_domain_token(tokens[1])
        if len(combo) >= 5 and combo in d:
            return True
            
    return False

async def fallback_search_website_fi(page, company_name):
    """
    Fallback tìm kiếm website qua DuckDuckGo / Bing TRỰC QUAN TRÊN TRÌNH DUYỆT (page)
    Người dùng quan sát trực tiếp từ khóa tìm kiếm và kết quả trên màn hình Chrome.
    Từ khóa tra cứu chỉ tra đúng tên công ty, không thêm bớt bất kỳ từ nào.
    Chỉ chấp nhận WEBSITE RIÊNG của doanh nghiệp, loại trừ toàn bộ trang nền tảng/danh bạ/mạng xã hội.
    """
    query = company_name.strip()
    q_clean = clean_company_name_fi(company_name)
    
    # 1. Tìm kiếm trên DuckDuckGo (trực quan trên Chrome)
    try:
        ddg_url = f"https://duckduckgo.com/?q={urllib.parse.quote(query)}"
        print(f"    -> Tìm kiếm trên DuckDuckGo: {query}", flush=True)
        await page.goto(ddg_url, wait_until="domcontentloaded", timeout=15000)
        await asyncio.sleep(2)
        await check_for_captcha(page, "DuckDuckGo Search")
        
        results = page.locator('article[data-testid="result"], div.result, [data-nrn="result"]')
        count = await results.count()
        if count > 0:
            for i in range(min(count, 6)):
                title_el = results.nth(i).locator('h2 a, a[data-testid="result-title-a"], .result__title a')
                t_text = await title_el.first.text_content() if await title_el.count() > 0 else ""
                href = await title_el.first.get_attribute('href') if await title_el.count() > 0 else ""
                if href and href.startswith('http'):
                    domain = urllib.parse.urlparse(href).netloc.lower().replace('www.', '')
                    
                    # BẮT BUỘC: Không được là trang nền tảng chung và phải chứa tên thương hiệu
                    if not is_strictly_company_domain_fi(domain, company_name):
                        continue
                        
                    matched, _ = is_valid_name_match_fi(q_clean, t_text)
                    if matched or any(normalize_fi_domain_token(t) in domain for t in q_clean.split() if len(t) >= 3):
                        print(f"    [+] Tìm thấy Website riêng doanh nghiệp: https://{domain}", flush=True)
                        return f"https://{domain}"
    except Exception as e:
        pass

    # 2. Fallback sang Bing trên trình duyệt nếu DuckDuckGo không có
    try:
        b_url = f"https://www.bing.com/search?q={urllib.parse.quote(query)}"
        print(f"    -> Tìm kiếm trên Bing: {query}", flush=True)
        await page.goto(b_url, wait_until="domcontentloaded", timeout=15000)
        await asyncio.sleep(2)
        await check_for_captcha(page, "Bing Search")
        
        try:
            accept_btn = page.locator('#bnp_btn_accept, button#bnp_btn_accept, button:has-text("Hyväksy"), button:has-text("Accept")')
            if await accept_btn.count() > 0:
                await accept_btn.first.click()
                await asyncio.sleep(1)
        except Exception:
            pass
            
        results = page.locator('li.b_algo')
        count = await results.count()
        if count > 0:
            for i in range(min(count, 6)):
                h2_a = results.nth(i).locator('h2 a')
                t_text = await h2_a.text_content() if await h2_a.count() > 0 else ""
                href = await h2_a.get_attribute('href') if await h2_a.count() > 0 else ""
                if not href:
                    cite = results.nth(i).locator('cite')
                    cite_text = await cite.text_content() if await cite.count() > 0 else ""
                    match = re.search(r'https?://[^\s›/]+', cite_text)
                    if match:
                        href = match.group(0)
                        
                if href and href.startswith('http'):
                    domain = urllib.parse.urlparse(href).netloc.lower().replace('www.', '')
                    
                    # BẮT BUỘC: Không được là trang nền tảng chung và phải chứa tên thương hiệu
                    if not is_strictly_company_domain_fi(domain, company_name):
                        continue
                        
                    matched, _ = is_valid_name_match_fi(q_clean, t_text)
                    if matched or any(normalize_fi_domain_token(t) in domain for t in q_clean.split() if len(t) >= 3):
                        print(f"    [+] Tìm thấy Website riêng doanh nghiệp: https://{domain}", flush=True)
                        return f"https://{domain}"
    except Exception:
        pass

    return ""

async def scrape_website_details_fi(http_session, raw_url, page=None):
    """
    Quét trang web tối ưu: Có Domain Cache, phát hiện trang liên hệ thực tế trên menu,
    hỗ trợ email chống bot Phần Lan (at / ät), decode HTML entities, và fallback Playwright render JS.
    """
    url = urllib.parse.unquote(raw_url.strip())
    if not url.startswith(('http://', 'https://')):
        url = 'https://' + url
        
    parsed = urllib.parse.urlparse(url)
    domain = parsed.netloc.lower().replace('www.', '')
    
    # 1. Kiểm tra Domain Cache trước
    if domain in DOMAIN_CACHE:
        return DOMAIN_CACHE[domain]
        
    collected_emails = set()
    collected_phones = []
    fb_url = ""
    phone_regex = re.compile(r'(?:\+358\s?|0)[1-9]\d{1,2}[\s-]?\d{3,4}[\s-]?\d{2,4}')
    
    # 2. Quét trang chủ
    homepage_html = ""
    try:
        res = await http_session.get(url, timeout=8, allow_redirects=True)
        if res.status_code == 200:
            homepage_html = res.text
            collected_emails.update(extract_emails_from_html(homepage_html))
            fb_url = extract_facebook_url(homepage_html)
            for ph in phone_regex.findall(homepage_html):
                if len(re.sub(r'\D', '', ph)) >= 8:
                    collected_phones.append(ph.strip())
    except Exception:
        if url.startswith('https://'):
            try:
                res = await http_session.get('http://' + url[8:], timeout=8, allow_redirects=True)
                if res.status_code == 200:
                    homepage_html = res.text
                    collected_emails.update(extract_emails_from_html(homepage_html))
                    fb_url = extract_facebook_url(homepage_html)
                    for ph in phone_regex.findall(homepage_html):
                        if len(re.sub(r'\D', '', ph)) >= 8:
                            collected_phones.append(ph.strip())
            except Exception:
                pass

    # 3. Quét thẻ <a> trên trang chủ tìm các trang liên hệ THỰC TẾ TRÊN MENU
    discovered_contact_urls = []
    if homepage_html:
        soup = BeautifulSoup(homepage_html, 'html.parser')
        keywords_contact = [
            'ota yhteyt', 'yhteystiedot', 'yhteys', 'asiakaspalvelu', 
            'rekrytoijat', 'tiimi', 'contact', 'meist', 'about', 
            'tietoa', 'henkilost', 'henkilöst', 'ihmiset', 'toimisto'
        ]
        
        for a in soup.find_all('a', href=True):
            href = urllib.parse.unquote(a['href'].strip())
            txt = a.get_text(strip=True).lower()
            href_l = href.lower()
            
            # Loại bỏ anchor nội bộ, javascript:, tel:, mailto: hoặc file tải về
            if href.startswith(('#', 'javascript:', 'tel:', 'mailto:')):
                continue
            if any(href_l.endswith(ext) for ext in ['.pdf', '.jpg', '.png', '.zip', '.docx']):
                continue
            # Loại bỏ các trang báo cáo tài chính/informaatio/blog
            if any(bad in href_l for bad in ['informaatio', 'sijoittaj', 'raport', 'taloustiet', 'blog/']):
                continue
                
            is_match = any(k in txt for k in keywords_contact) or any(k in href_l for k in keywords_contact)
            if is_match:
                full_u = urllib.parse.urljoin(url, href)
                parsed_u = urllib.parse.urlparse(full_u)
                if parsed_u.netloc.lower().replace('www.', '') == domain and full_u != url:
                    if full_u not in discovered_contact_urls:
                        discovered_contact_urls.append(full_u)
                        if len(discovered_contact_urls) >= 8:
                            break

    # Ưu tiên các trang liên hệ thực tế tìm thấy trên menu
    contact_urls = discovered_contact_urls
    # Nếu không tìm thấy link nào trên menu, mới thử các slug mặc định
    if not contact_urls:
        for slug in ['/yhteystiedot/', '/ota-yhteytta/', '/meista/', '/yhteystiedot', '/ota-yhteytta']:
            contact_urls.append(urllib.parse.urljoin(url, slug))

    # 4. Quét tối đa 5 trang liên hệ tốt nhất
    for cu in contact_urls[:5]:
        try:
            c_res = await http_session.get(cu, timeout=7, allow_redirects=True)
            if c_res.status_code == 200:
                collected_emails.update(extract_emails_from_html(c_res.text))
                for ph in phone_regex.findall(c_res.text):
                    if len(re.sub(r'\D', '', ph)) >= 8:
                        collected_phones.append(ph.strip())
                if not fb_url:
                    fb_url = extract_facebook_url(c_res.text)
        except Exception:
            continue

    # 5. Fallback Playwright (render JavaScript): Nếu vẫn chưa có email và có browser page
    if not collected_emails and page:
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=10000)
            await asyncio.sleep(2)
            rendered_html = await page.content()
            collected_emails.update(extract_emails_from_html(rendered_html))
            if not fb_url:
                fb_url = extract_facebook_url(rendered_html)
        except Exception:
            pass

    emails_str = prioritize_emails(collected_emails, domain)
    phone_str = collected_phones[0] if collected_phones else ""
    DOMAIN_CACHE[domain] = (emails_str, phone_str, fb_url)
    return emails_str, phone_str, fb_url

async def scrape_facebook_email_fi(page, fb_url):
    """
    Quét Fanpage Facebook với FB_CACHE và timeout nhanh
    """
    if not fb_url:
        return ""
    clean_fb = fb_url.rstrip('/').lower()
    if clean_fb in FB_CACHE:
        return FB_CACHE[clean_fb]
        
    email_found = ""
    try:
        about_url = fb_url.rstrip('/') + '/about'
        await page.goto(about_url, wait_until="domcontentloaded", timeout=10000)
        await asyncio.sleep(2)
        
        try:
            close_btn = page.locator('div[aria-label="Sulje"], div[aria-label="Close"], i[data-visualcompletion="css-img"]')
            if await close_btn.count() > 0:
                await close_btn.first.click()
                await asyncio.sleep(1)
        except Exception:
            pass
            
        content = await page.content()
        emails = extract_emails_from_html(content)
        valid = [e for e in emails if 'facebook' not in e and 'fb.com' not in e]
        if valid:
            email_found = valid[0]
            FB_CACHE[clean_fb] = email_found
            return email_found
            
        # Fallback trang chính
        await page.goto(fb_url, wait_until="domcontentloaded", timeout=8000)
        await asyncio.sleep(1.5)
        content2 = await page.content()
        emails2 = extract_emails_from_html(content2)
        valid2 = [e for e in emails2 if 'facebook' not in e and 'fb.com' not in e]
        if valid2:
            email_found = valid2[0]
    except Exception:
        pass
        
    FB_CACHE[clean_fb] = email_found
    return email_found

def sync_to_all_branches(updated_rows):
    """Đồng bộ các email/sđt/facebook mới tìm được sang file tất cả chi nhánh"""
    all_branches_csv = os.path.join(BASE_DIR, "agency phần lan - tất cả chi nhánh.csv")
    if not os.path.exists(all_branches_csv):
        return
    try:
        enrich_map_bid = {}
        enrich_map_name = {}
        for r in updated_rows:
            bid = r.get('business_id', '').strip()
            name = clean_company_name_fi(r.get('name', ''))
            em = r.get('email', '').strip()
            ph = r.get('phone', '').strip()
            fb = r.get('facebook_url', '').strip()
            if em or ph or fb:
                if bid:
                    enrich_map_bid[bid] = (em, ph, fb)
                if name:
                    enrich_map_name[name] = (em, ph, fb)
                    
        with open(all_branches_csv, 'r', encoding='utf-8-sig') as f:
            b_reader = csv.DictReader(f)
            b_fields = list(b_reader.fieldnames)
            b_rows = list(b_reader)
            
        modified = False
        for br in b_rows:
            bid = br.get('business_id', '').strip()
            name = clean_company_name_fi(br.get('name', ''))
            match = enrich_map_bid.get(bid) or enrich_map_name.get(name)
            if match:
                em, ph, fb = match
                if em and not br.get('email', '').strip():
                    br['email'] = em
                    modified = True
                if ph and not br.get('phone', '').strip():
                    br['phone'] = ph
                    modified = True
                if fb and not br.get('facebook_url', '').strip():
                    br['facebook_url'] = fb
                    modified = True
                    
        if modified:
            with open(all_branches_csv, 'w', encoding='utf-8-sig', newline='') as f:
                writer = csv.DictWriter(f, fieldnames=b_fields)
                writer.writeheader()
                writer.writerows(b_rows)
    except Exception:
        pass

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def save_cache(cache):
    os.makedirs(os.path.dirname(CACHE_FILE), exist_ok=True)
    with open(CACHE_FILE, 'w', encoding='utf-8') as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

def save_outputs(rows, fieldnames):
    new_cols = ['tag_doanh_nghiep_maps', 'facebook_url', 'dia_chi_maps', 'nguon_enrich']
    for c in new_cols:
        if c not in fieldnames:
            fieldnames.append(c)
            
    # Save CSV UTF-8-BOM với fallback backup nếu file đang mở trong Excel
    try:
        with open(OUTPUT_CSV, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
    except PermissionError:
        backup_csv = OUTPUT_CSV.replace('.csv', '_backup.csv')
        try:
            with open(backup_csv, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(rows)
            print(f"  [!] File CSV đang mở, đã lưu động vào: {os.path.basename(backup_csv)}", flush=True)
        except Exception:
            pass
    except Exception as e:
        print(f"  [!] Lỗi lưu CSV: {e}", flush=True)
        
    # Đồng bộ sang file tất cả chi nhánh
    sync_to_all_branches(rows)
        
    # Save Excel: bảo toàn các sheet hiện có
    try:
        if os.path.exists(OUTPUT_XLSX):
            try:
                wb = openpyxl.load_workbook(OUTPUT_XLSX)
            except Exception:
                wb = openpyxl.Workbook()
        else:
            wb = openpyxl.Workbook()
            
        sheet_name = "Doanh nghiệp (Lọc trùng)"
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            ws.delete_rows(1, ws.max_row + 1)
        else:
            ws = wb.create_sheet(title=sheet_name, index=0)
            
        ws.append(fieldnames)
        for r in rows:
            ws.append([r.get(f, '') for f in fieldnames])
        wb.save(OUTPUT_XLSX)
    except PermissionError:
        backup_xlsx = OUTPUT_XLSX.replace('.xlsx', '_backup.xlsx')
        try:
            wb.save(backup_xlsx)
            print(f"  [!] File Excel đang mở, đã lưu động vào: {os.path.basename(backup_xlsx)}", flush=True)
        except Exception:
            pass
    except Exception as e:
        print(f"  [!] Lỗi lưu Excel: {e}", flush=True)

async def main():
    is_web_only = '--web-only' in sys.argv
    print("=" * 75, flush=True)
    if is_web_only:
        print(" CHẾ ĐỘ ƯU TIÊN: CÀO EMAIL CHO CÁC DOANH NGHIỆP CÓ SẴN WEBSITE ", flush=True)
        print(" (Tối ưu tốc độ cao: Bỏ qua Maps/Search, tập trung quét Web + Facebook) ", flush=True)
    else:
        print(" HỆ THỐNG LÀM GIÀU DỮ LIỆU AGENCY PHẦN LAN (TỐI ƯU HÓA HIỆU NĂNG) ", flush=True)
        print(" Chế độ quan sát trực quan: HEADLESS = FALSE (Mở cửa sổ Chrome) ", flush=True)
    print("=" * 75, flush=True)
    
    if not os.path.exists(INPUT_CSV):
        print(f"[-] Không tìm thấy file: {INPUT_CSV}", flush=True)
        return
        
    rows = []
    fieldnames = []
    with open(INPUT_CSV, 'r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        fieldnames = list(reader.fieldnames)
        rows = list(reader)
        
    # Lọc danh sách cần làm giàu email
    if is_web_only:
        target_rows = [r for r in rows if r.get('website', '').strip() and not r.get('email', '').strip()]
        print(f"[+] Tổng số agency duy nhất: {len(rows)} | Có website cần cào email: {len(target_rows)} agency", flush=True)
    else:
        target_rows = [r for r in rows if not r.get('email', '').strip()]
        # Ưu tiên các dòng có sẵn website chạy trước
        target_rows.sort(key=lambda r: 0 if r.get('website', '').strip() else 1)
        has_web_count = sum(1 for r in target_rows if r.get('website', '').strip())
        print(f"[+] Tổng số agency duy nhất: {len(rows)} | Cần làm giàu email: {len(target_rows)} agency", flush=True)
        print(f"    -> Trong đó có {has_web_count} agency có sẵn website (được ưu tiên quét trước)", flush=True)
    
    cache = load_cache()
    print(f"[+] Đã có trong Cache: {len(cache)} agency đã xử lý trước đó", flush=True)
    
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False, channel="chrome")
        context = await browser.new_context(
            locale="fi-FI",
            viewport={"width": 1280, "height": 850},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
        )
        page = await context.new_page()
        
        async with AsyncSession(impersonate='chrome124', verify=False) as http_session:
            processed_count = 0
            try:
                for idx, r in enumerate(target_rows, 1):
                    c_name = r['name'].strip()
                    cache_key = c_name.lower()
                    
                    # 1. Kiểm tra cache
                    if cache_key in cache:
                        c_data = cache[cache_key]
                        if c_data.get('email'):
                            r['email'] = c_data['email']
                        if c_data.get('phone') and not r.get('phone'):
                            r['phone'] = c_data['phone']
                        if c_data.get('website') and not r.get('website'):
                            r['website'] = c_data['website']
                        if c_data.get('tag_maps'):
                            r['tag_doanh_nghiep_maps'] = c_data['tag_maps']
                        if c_data.get('facebook_url'):
                            r['facebook_url'] = c_data['facebook_url']
                        if c_data.get('address'):
                            r['dia_chi_maps'] = c_data['address']
                        continue
                        
                    processed_count += 1
                    print(f"\n[{processed_count}/{len(target_rows)}] Xử lý: {c_name}", flush=True)
                    
                    cur_phone = r.get('phone', '').strip()
                    cur_web = r.get('website', '').strip()
                    cur_email = r.get('email', '').strip()
                    tag_maps = ""
                    addr_maps = ""
                    fb_url = ""
                    source = []
                    
                    # 2. CHIẾN LƯỢC TỐI ƯU: NẾU ĐÃ CÓ WEBSITE TỪ TRƯỚC -> QUÉT WEBSITE NGAY
                    if cur_web:
                        print(f"  -> Quét nhanh Website sẵn có ({cur_web})...", flush=True)
                        web_email, found_phone, found_fb = await scrape_website_details_fi(http_session, cur_web, page=page)
                        if web_email:
                            cur_email = web_email
                            source.append("Website")
                            print(f"    [+] EMAIL TỪ WEB: {cur_email}", flush=True)
                        if found_phone and not cur_phone:
                            cur_phone = found_phone
                            print(f"    [+] SĐT TỪ WEB: {cur_phone}", flush=True)
                        if found_fb:
                            fb_url = found_fb
                            print(f"    [+] LINK FACEBOOK: {fb_url}", flush=True)
                            
                    # 3. NẾU CHƯA CÓ WEBSITE HOẶC CẦN TÌM THÊM THÔNG TIN LIÊN HỆ -> MỚI TRA MAPS (Chỉ chạy khi không ở chế độ --web-only)
                    if not is_web_only and (not cur_web or not cur_phone or not cur_email):
                        print(f"  -> Tra cứu trên Google Maps (hl=fi)...", flush=True)
                        maps_info = await search_google_maps_fi(page, c_name)
                        if maps_info["matched"]:
                            source.append("GoogleMaps")
                            tag_maps = maps_info["tag_maps"]
                            addr_maps = maps_info["address"]
                            print(f"    [Maps Khớp ({maps_info['match_reason']})]: {maps_info['maps_title']}", flush=True)
                            if tag_maps:
                                print(f"      - Tag ngành (FI): {tag_maps}", flush=True)
                            if maps_info["phone"] and not cur_phone:
                                cur_phone = maps_info["phone"]
                                print(f"      - SĐT mới (Maps): {cur_phone}", flush=True)
                            if maps_info["website"] and not cur_web:
                                m_web = maps_info["website"].strip()
                                m_dom = urllib.parse.urlparse(m_web).netloc.lower().replace('www.', '')
                                if 'facebook.com' in m_dom or 'fb.com' in m_dom:
                                    if not fb_url:
                                        fb_url = m_web
                                        print(f"      - Phát hiện Facebook từ Maps: {fb_url}", flush=True)
                                elif is_strictly_company_domain_fi(m_dom, c_name):
                                    cur_web = m_web
                                    print(f"      - Website mới (Maps): {cur_web}", flush=True)
                                    # Cào website mới phát hiện từ Maps
                                    web_email, found_phone, found_fb = await scrape_website_details_fi(http_session, cur_web, page=page)
                                    if web_email and not cur_email:
                                        cur_email = web_email
                                        source.append("Website(Maps)")
                                        print(f"      [+] EMAIL MỚI: {cur_email}", flush=True)
                                    if found_phone and not cur_phone:
                                        cur_phone = found_phone
                                    if found_fb and not fb_url:
                                        fb_url = found_fb
                        else:
                            print(f"    [Maps Không khớp/Không thấy]", flush=True)
                            
                    # 4. BƯỚC FALLBACK SEARCH BING / DUCKDUCKGO NẾU VẪN CHƯA CÓ WEBSITE (Chỉ chạy khi không ở chế độ --web-only)
                    if not is_web_only and not cur_web:
                        print(f"  -> Fallback Search (Bing/DDG) tìm website...", flush=True)
                        fb_web = await fallback_search_website_fi(page, c_name)
                        if fb_web:
                            cur_web = fb_web
                            source.append("SearchFallback")
                            print(f"    - Website tìm thấy: {cur_web}", flush=True)
                            web_email, found_phone, found_fb = await scrape_website_details_fi(http_session, cur_web, page=page)
                            if web_email and not cur_email:
                                cur_email = web_email
                                source.append("Website(Search)")
                                print(f"    [+] EMAIL WEBSITE: {cur_email}", flush=True)
                            if found_phone and not cur_phone:
                                cur_phone = found_phone
                            if found_fb and not fb_url:
                                fb_url = found_fb
                        else:
                            print(f"    - Không tìm thấy website trên search", flush=True)
                            
                    # 5. QUÉT FACEBOOK NẾU VẪN CHƯA CÓ EMAIL MÀ CÓ LINK FB
                    if not cur_email and fb_url:
                        print(f"  -> Quét Fanpage Facebook ({fb_url}) tìm email...", flush=True)
                        fb_email = await scrape_facebook_email_fi(page, fb_url)
                        if fb_email:
                            cur_email = fb_email
                            source.append("Facebook")
                            print(f"    [+] EMAIL FACEBOOK: {cur_email}", flush=True)
                        else:
                            print(f"    - Không thấy email trên Facebook", flush=True)
                            
                    # CẬP NHẬT DỮ LIỆU
                    r['email'] = cur_email
                    if cur_phone:
                        r['phone'] = cur_phone
                    if cur_web:
                        r['website'] = cur_web
                    if tag_maps:
                        r['tag_doanh_nghiep_maps'] = tag_maps
                    if addr_maps:
                        r['dia_chi_maps'] = addr_maps
                    if fb_url:
                        r['facebook_url'] = fb_url
                    r['nguon_enrich'] = ", ".join(source) if source else "None"
                    
                    # Lưu cache ngay lập tức sau mỗi công ty
                    cache[cache_key] = {
                        "email": cur_email,
                        "phone": cur_phone,
                        "website": cur_web,
                        "tag_maps": tag_maps,
                        "address": addr_maps,
                        "facebook_url": fb_url,
                        "source": r['nguon_enrich']
                    }
                    save_cache(cache)
                    
                    # Lưu CSV / Excel sau mỗi 3 công ty hoặc khi có email mới
                    if processed_count % 3 == 0 or cur_email:
                        save_outputs(rows, fieldnames)
                        
            except (KeyboardInterrupt, asyncio.CancelledError):
                print("\n[!] Đã nhận tín hiệu dừng (Ctrl+C). Đang lưu toàn bộ tiến trình...", flush=True)
            except Exception as e:
                print(f"\n[!] Lỗi trong quá trình chạy: {e}", flush=True)
            finally:
                save_cache(cache)
                save_outputs(rows, fieldnames)
                print("[+] Đã lưu file an toàn trước khi đóng trình duyệt!", flush=True)
                try:
                    await browser.close()
                except Exception:
                    pass
                    
    print("\n" + "=" * 75, flush=True)
    print("                 HOÀN TẤT LÀM GIÀU DỮ LIỆU AGENCY PHẦN LAN!", flush=True)
    print("=" * 75, flush=True)
    print(f"- File CSV kết quả: {OUTPUT_CSV}", flush=True)
    print(f"- File Excel kết quả: {OUTPUT_XLSX}", flush=True)
    print("=" * 75, flush=True)

if __name__ == '__main__':
    asyncio.run(main())
