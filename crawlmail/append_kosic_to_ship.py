# -*- coding: utf-8 -*-
"""
Tích hợp danh sách 65 Xưởng đóng tàu thuộc Hiệp hội Hợp tác xã Công nghiệp Đóng tàu Hàn Quốc (KOSIC)
từ http://www.kosic.or.kr/ vào tệp KR_SHIP_HIPC_0RAW.csv.
"""

import os
import sys
import csv
import openpyxl
import pandas as pd
from pathlib import Path

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

ROOT_DIR = Path(__file__).resolve().parent.parent
TARGET_CSV = ROOT_DIR / "KR_SHIP_HIPC_0RAW.csv"
KOSIC_EXCEL = ROOT_DIR / "kosic_members.xlsx"

HEADERS = [
    'No.', 'Cong ty', 'Chuc danh', 'Nguoi lien he', 'SDT', 
    'Lien He', 'Email', 'Lien He mail', 'Dia chi', 'Luong', 
    'Ngay dang', 'Han tuyen', 'Check gui', 'Last Subject', 
    'Last Body HTML', 'Trang thai Reply', 'Lan Follow-up', 
    'Ngay Follow-up gan nhat', 'Mailbox da dung', 'Category'
]

def clean_phone(p):
    if not p:
        return ""
    p_str = str(p).strip()
    if not p_str or p_str.lower() == 'none' or p_str.lower() == 'nan':
        return ""
    if not p_str.startswith("'"):
        return f"'{p_str}"
    return p_str

def main():
    if not TARGET_CSV.exists():
        print(f"[!] Không tìm thấy tệp đích: {TARGET_CSV}")
        return

    # 1. Đọc dữ liệu hiện có trong KR_SHIP_HIPC_0RAW.csv
    df_existing = pd.read_csv(TARGET_CSV, dtype=str).fillna("")
    print(f"Số lượng công ty hiện có trong KR_SHIP_HIPC_0RAW.csv: {len(df_existing)}")
    
    existing_companies = set(df_existing["Cong ty"].str.strip().tolist())

    # 2. Đọc dữ liệu từ file Excel KOSIC
    wb = openpyxl.load_workbook(KOSIC_EXCEL, data_only=True)
    sheet = wb.active

    new_rows = []
    updated_count = 0

    for r in range(2, sheet.max_row + 1):
        cname = sheet.cell(r, 2).value
        if not cname:
            continue
        cname = str(cname).strip()
        
        rep = str(sheet.cell(r, 3).value or "").strip()
        web = str(sheet.cell(r, 4).value or "").strip()
        addr = str(sheet.cell(r, 5).value or "").strip()
        region = str(sheet.cell(r, 6).value or "").strip()
        phone = clean_phone(sheet.cell(r, 7).value)
        fax = str(sheet.cell(r, 8).value or "").strip()
        hull = str(sheet.cell(r, 9).value or "").strip()
        ship_types = str(sheet.cell(r, 10).value or "").strip()

        # Format chức danh & phân loại
        title_desc = []
        if hull:
            title_desc.append(f"Vỏ tàu: {hull}")
        if ship_types:
            title_desc.append(f"Loại tàu: {ship_types}")
        chuc_danh = " | ".join(title_desc) if title_desc else "Đóng mới & sửa chữa tàu biển"

        contact_web = web if web and web.startswith("http") else "http://www.kosic.or.kr/"
        cat = f"한국중소조선공업협동조합 (KOSIC - {region})" if region else "한국중소조선공업협동조합 (KOSIC)"

        # Kiểm tra nếu đã có trong danh sách hiện tại
        if cname in existing_companies:
            # Cập nhật thông tin chi tiết vào dòng hiện có
            idx = df_existing[df_existing["Cong ty"] == cname].index
            if len(idx) > 0:
                i = idx[0]
                if phone and not df_existing.at[i, "SDT"]:
                    df_existing.at[i, "SDT"] = phone
                if web and not df_existing.at[i, "Lien He"]:
                    df_existing.at[i, "Lien He"] = web
                if addr and len(addr) > len(df_existing.at[i, "Dia chi"]):
                    df_existing.at[i, "Dia chi"] = addr
                updated_count += 1
        else:
            new_row = {
                "No.": "",
                "Cong ty": cname,
                "Chuc danh": chuc_danh,
                "Nguoi lien he": rep,
                "SDT": phone,
                "Lien He": contact_web,
                "Email": "",
                "Lien He mail": contact_web,
                "Dia chi": addr,
                "Luong": "",
                "Ngay dang": "2026-10-06",
                "Han tuyen": "",
                "Check gui": "",
                "Last Subject": "",
                "Last Body HTML": "",
                "Trang thai Reply": "",
                "Lan Follow-up": "",
                "Ngay Follow-up gan nhat": "",
                "Mailbox da dung": "",
                "Category": cat
            }
            new_rows.append(new_row)

    print(f" -> Cập nhật thông tin chi tiết: {updated_count} công ty trùng lặp")
    print(f" -> Bổ sung mới: {len(new_rows)} xưởng đóng tàu KOSIC")

    if new_rows:
        df_new = pd.DataFrame(new_rows)
        for col in HEADERS:
            if col not in df_new.columns:
                df_new[col] = ""
        df_new = df_new[HEADERS]
        df_combined = pd.concat([df_existing, df_new], ignore_index=True)
    else:
        df_combined = df_existing

    # Đánh số lại No. từ 1 đến N
    df_combined["No."] = range(1, len(df_combined) + 1)
    
    # Đảm bảo SĐT luôn có nháy đơn
    df_combined["SDT"] = df_combined["SDT"].apply(clean_phone)

    df_combined.to_csv(TARGET_CSV, index=False, encoding="utf-8-sig")
    print(f"[V] Đã lưu thành công tổng cộng {len(df_combined)} doanh nghiệp vào {TARGET_CSV}")

    # Xóa file Excel tạm
    if KOSIC_EXCEL.exists():
        try:
            wb.close()
            os.remove(KOSIC_EXCEL)
            print(" -> Đã xóa tệp tạm kosic_members.xlsx")
        except Exception:
            pass

if __name__ == "__main__":
    main()
