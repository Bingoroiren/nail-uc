import sys
import io
import json
import time
import os
import csv
from datetime import datetime
from curl_cffi import requests

# Đảm bảo in tiếng Việt trên console Windows không bị lỗi font
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', line_buffering=True)

BASE_SECTIONS_URL = "https://webgate.ec.europa.eu/tracesnt/directory/listing/establishment/publication?classificationSectionChapter=food&countryCode=LT"
ESTABLISHMENTS_API_TEMPLATE = "https://webgate.ec.europa.eu/tracesnt/directory/listing/establishment/publication/establishments/LT/{code}?max=1000&offset=0&sort=name.value"

OUTPUT_CSV_ALL = "chế biến thực phẩm litva.csv"
OUTPUT_CSV_UNIQUE = "chế biến thực phẩm litva (danh sách công ty duy nhất).csv"
OUTPUT_XLSX = "chế biến thực phẩm litva.xlsx"
CACHE_JSON = "crawlmail/cache_che_bien_thuc_pham_litva.json"

def crawl_traces_food_lithuania():
    print("=" * 70)
    print("   BỘ CÀO DANH SÁCH DOANH NGHIỆP CHẾ BIẾN THỰC PHẨM LITVA (TRACES NT - EU)")
    print("   Nguồn: https://webgate.ec.europa.eu/tracesnt/...")
    print("=" * 70)

    session = requests.Session()
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "en-US,en;q=0.9,vi;q=0.8",
    }

    # 1. Lấy danh sách các phân ngành thực phẩm (Food Sections) của Lithuania
    print("[*] Đang tải danh sách 14 phân ngành chế biến thực phẩm tại Litva...")
    sections = []
    for attempt in range(3):
        try:
            r = session.get(BASE_SECTIONS_URL, headers=headers, impersonate="chrome124", timeout=25)
            if r.status_code == 200:
                sections = r.json()
                print(f"[+] Tìm thấy {len(sections)} phân ngành thực phẩm.")
                break
        except Exception as e:
            print(f"[!] Lỗi kết nối (thử lại {attempt+1}/3): {e}")
            time.sleep(2)
    else:
        print("[!] Không thể tải danh sách phân ngành từ TRACES NT.")
        return

    # 2. Duyệt qua từng phân ngành (tương đương nhấn vào biểu tượng con mắt trên web)
    all_records = []
    
    print("\n[*] Đang cào chi tiết từng phân ngành (nhấn vào con mắt)...")
    for idx, sec in enumerate(sections, 1):
        sec_info = sec.get('establishmentListingId', {}).get('classificationSectionId', {})
        code = sec_info.get('code')
        sec_id = sec_info.get('id', '')
        sec_name = sec_info.get('translation', '')
        expected_count = sec.get('numberOfEstablishments', 0)

        url = ESTABLISHMENTS_API_TEMPLATE.format(code=code)
        view_web_url = f"https://webgate.ec.europa.eu/tracesnt/directory/listing/establishment/publication/index#!/view/LT/{sec_id}"

        print(f"[{idx:>2}/{len(sections)}] Ngành '{code}' - {sec_name[:40]:<40} (Dự kiến: {expected_count} cơ sở)...", end=" ", flush=True)

        items = []
        for attempt in range(4):
            try:
                res = session.get(url, headers=headers, impersonate="chrome124", timeout=30)
                if res.status_code == 200:
                    items = res.json()
                    break
                else:
                    time.sleep(1.5)
            except Exception:
                time.sleep(2)

        print(f"-> Thu thập thành công {len(items)} cơ sở.")

        for it in items:
            name = it.get('operatorName') or ''
            approval_no = it.get('approvalNumber') or ''

            # Địa chỉ
            addr = it.get('address') or {}
            street_val = addr.get('street', {})
            street = street_val.get('value') if isinstance(street_val, dict) else (it.get('street') or '')
            
            city_ref = addr.get('cityReference') or it.get('cityReference') or {}
            city = city_ref.get('name') or ''
            postal_code = city_ref.get('postalCode') or ''
            
            regions = city_ref.get('hierarchicalRegions') or []
            region_name = regions[0].get('name') if regions else ''

            # Hoạt động (Activities)
            acts = []
            for act in it.get('operatorActivityTypes') or []:
                code_trans = act.get('officialCodeAndTranslation') or act.get('translation') or act.get('officialCode')
                if code_trans and code_trans not in acts:
                    acts.append(code_trans)
            activities_str = "; ".join(acts)

            # Ghi chú
            remarks = "; ".join([str(rem.get('value') or rem) for rem in (it.get('remarks') or []) if rem])

            # Ngày cập nhật
            last_update_raw = it.get('lastUpdate') or ''
            last_update = last_update_raw[:10] if len(last_update_raw) >= 10 else last_update_raw

            record = {
                "ten_cong_ty": name.strip(),
                "ma_phe_duyet": approval_no.strip(),
                "ma_nganh": code,
                "ten_nganh": sec_name,
                "hoat_dong": activities_str,
                "dia_chi": street.strip(),
                "thanh_pho": city.strip(),
                "ma_buu_chinh": postal_code.strip(),
                "vung_tinh": region_name.strip(),
                "quoc_gia": "Lithuania",
                "ngay_cap_nhat": last_update,
                "ghi_chu": remarks,
                "traces_url": view_web_url
            }
            all_records.append(record)

        time.sleep(0.5)

    # 3. Lưu dữ liệu cache JSON
    os.makedirs(os.path.dirname(CACHE_JSON) if os.path.dirname(CACHE_JSON) else '.', exist_ok=True)
    with open(CACHE_JSON, 'w', encoding='utf-8') as f:
        json.dump(all_records, f, ensure_ascii=False, indent=2)
    print(f"\n[+] Đã lưu dữ liệu thô vào: {CACHE_JSON}")

    # 4. Xuất file CSV 1: Toàn bộ danh sách (tất cả các bản ghi phân theo ngành)
    fieldnames_all = [
        "stt",
        "ten_cong_ty",
        "ma_phe_duyet",
        "ma_nganh",
        "ten_nganh",
        "hoat_dong",
        "dia_chi",
        "thanh_pho",
        "ma_buu_chinh",
        "vung_tinh",
        "quoc_gia",
        "ngay_cap_nhat",
        "ghi_chu",
        "traces_url"
    ]

    with open(OUTPUT_CSV_ALL, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames_all)
        writer.writeheader()
        for idx, rec in enumerate(all_records, 1):
            row = {"stt": idx}
            row.update(rec)
            writer.writerow(row)
    print(f"[+] Đã xuất file CSV đầy đủ (từng ngành): {OUTPUT_CSV_ALL}")

    # 5. Xuất file CSV 2: Danh sách công ty duy nhất (gộp các ngành công ty hoạt động)
    unique_companies = {}
    for rec in all_records:
        key = rec['ten_cong_ty'].upper().strip()
        if not key:
            key = rec['ma_phe_duyet']

        if key not in unique_companies:
            unique_companies[key] = {
                "ten_cong_ty": rec['ten_cong_ty'],
                "ma_phe_duyet": rec['ma_phe_duyet'],
                "cac_nganh": [f"{rec['ma_nganh']} ({rec['ten_nganh']})"],
                "cac_hoat_dong": [rec['hoat_dong']] if rec['hoat_dong'] else [],
                "dia_chi": rec['dia_chi'],
                "thanh_pho": rec['thanh_pho'],
                "ma_buu_chinh": rec['ma_buu_chinh'],
                "vung_tinh": rec['vung_tinh'],
                "quoc_gia": "Lithuania",
                "ngay_cap_nhat_moi_nhat": rec['ngay_cap_nhat'],
                "so_luong_nganh_cap_phep": 1
            }
        else:
            item = unique_companies[key]
            sec_entry = f"{rec['ma_nganh']} ({rec['ten_nganh']})"
            if sec_entry not in item['cac_nganh']:
                item['cac_nganh'].append(sec_entry)
            if rec['hoat_dong'] and rec['hoat_dong'] not in item['cac_hoat_dong']:
                item['cac_hoat_dong'].append(rec['hoat_dong'])
            if rec['ma_phe_duyet'] and rec['ma_phe_duyet'] not in item['ma_phe_duyet']:
                item['ma_phe_duyet'] += f"; {rec['ma_phe_duyet']}"
            item['so_luong_nganh_cap_phep'] = len(item['cac_nganh'])
            if rec['ngay_cap_nhat'] > item['ngay_cap_nhat_moi_nhat']:
                item['ngay_cap_nhat_moi_nhat'] = rec['ngay_cap_nhat']

    fieldnames_unique = [
        "stt",
        "ten_cong_ty",
        "ma_phe_duyet",
        "so_luong_nganh_cap_phep",
        "cac_nganh_hoat_dong",
        "cac_hoat_dong_chi_tiet",
        "dia_chi",
        "thanh_pho",
        "ma_buu_chinh",
        "vung_tinh",
        "quoc_gia",
        "ngay_cap_nhat_moi_nhat"
    ]

    unique_list = list(unique_companies.values())
    unique_list.sort(key=lambda x: x['ten_cong_ty'])

    with open(OUTPUT_CSV_UNIQUE, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames_unique)
        writer.writeheader()
        for idx, rec in enumerate(unique_list, 1):
            row = {
                "stt": idx,
                "ten_cong_ty": rec['ten_cong_ty'],
                "ma_phe_duyet": rec['ma_phe_duyet'],
                "so_luong_nganh_cap_phep": rec['so_luong_nganh_cap_phep'],
                "cac_nganh_hoat_dong": " | ".join(rec['cac_nganh']),
                "cac_hoat_dong_chi_tiet": " | ".join(rec['cac_hoat_dong']),
                "dia_chi": rec['dia_chi'],
                "thanh_pho": rec['thanh_pho'],
                "ma_buu_chinh": rec['ma_buu_chinh'],
                "vung_tinh": rec['vung_tinh'],
                "quoc_gia": rec['quoc_gia'],
                "ngay_cap_nhat_moi_nhat": rec['ngay_cap_nhat_moi_nhat']
            }
            writer.writerow(row)
    print(f"[+] Đã xuất file CSV công ty duy nhất: {OUTPUT_CSV_UNIQUE}")

    # 6. Xuất file Excel (.xlsx) 2 trang tính nếu có thư viện openpyxl
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        
        # Sheet 1: Tất cả cơ sở (923)
        ws1 = wb.active
        ws1.title = "Tất cả cơ sở theo ngành"
        ws1.views.sheetView[0].showGridLines = True

        headers1 = [
            "STT", "Tên Công Ty / Cơ Sở", "Mã Phê Duyệt", "Mã Ngành", "Tên Phân Ngành",
            "Hoạt Động Cấp Phép", "Địa Chỉ", "Thành Phố", "Mã Bưu Điện", "Vùng / Tỉnh",
            "Quốc Gia", "Ngày Cập Nhật", "Ghi Chú", "Link TRACES NT"
        ]
        ws1.append(headers1)

        for idx, rec in enumerate(all_records, 1):
            ws1.append([
                idx,
                rec['ten_cong_ty'],
                rec['ma_phe_duyet'],
                rec['ma_nganh'],
                rec['ten_nganh'],
                rec['hoat_dong'],
                rec['dia_chi'],
                rec['thanh_pho'],
                rec['ma_buu_chinh'],
                rec['vung_tinh'],
                rec['quoc_gia'],
                rec['ngay_cap_nhat'],
                rec['ghi_chu'],
                rec['traces_url']
            ])

        # Sheet 2: Công ty duy nhất (659)
        ws2 = wb.create_sheet(title="Doanh nghiệp duy nhất")
        ws2.views.sheetView[0].showGridLines = True

        headers2 = [
            "STT", "Tên Doanh Nghiệp", "Mã Phê Duyệt", "Số Ngành", "Các Phân Ngành Thực Phẩm",
            "Các Hoạt Động Chi Tiết", "Địa Chỉ", "Thành Phố", "Mã Bưu Điện", "Vùng / Tỉnh",
            "Quốc Gia", "Ngày Cập Nhật Mới Nhất"
        ]
        ws2.append(headers2)

        for idx, rec in enumerate(unique_list, 1):
            ws2.append([
                idx,
                rec['ten_cong_ty'],
                rec['ma_phe_duyet'],
                rec['so_luong_nganh_cap_phep'],
                " | ".join(rec['cac_nganh']),
                " | ".join(rec['cac_hoat_dong']),
                rec['dia_chi'],
                rec['thanh_pho'],
                rec['ma_buu_chinh'],
                rec['vung_tinh'],
                rec['quoc_gia'],
                rec['ngay_cap_nhat_moi_nhat']
            ])

        # Style header cho cả 2 sheet
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for ws in [ws1, ws2]:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Tự căn chỉnh độ rộng cột
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col[:100]:
                    if cell.value:
                        val_str = str(cell.value)
                        max_len = max(max_len, min(len(val_str), 50))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(OUTPUT_XLSX)
        print(f"[+] Đã xuất file Excel đẹp 2 sheet: {OUTPUT_XLSX}")
    except Exception as e:
        print(f"[!] Lỗi tạo file Excel: {e}")

    print("\n" + "=" * 70)
    print("                 HOÀN THÀNH CÀO TRACES NT - LITVA")
    print("=" * 70)
    print(f"- Tổng số lượt cấp phép cơ sở (theo ngành) : {len(all_records)}")
    print(f"- Tổng số doanh nghiệp thực phẩm duy nhất  : {len(unique_list)}")
    print(f"- File CSV kết quả chính                     : {os.path.abspath(OUTPUT_CSV_ALL)}")
    print(f"- File CSV doanh nghiệp duy nhất             : {os.path.abspath(OUTPUT_CSV_UNIQUE)}")
    if os.path.exists(OUTPUT_XLSX):
        print(f"- File Excel (.xlsx) 2 trang tính           : {os.path.abspath(OUTPUT_XLSX)}")
    print("=" * 70)

if __name__ == "__main__":
    crawl_traces_food_lithuania()
