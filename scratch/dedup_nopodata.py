import csv
import re
import sys
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

INPUT_FILE = 'NOPODATA.csv'
OUTPUT_CSV = 'NOPODATA (ĐÃ LỌC TRÙNG EMAIL).csv'
OUTPUT_XLSX = 'NOPODATA (ĐÃ LỌC TRÙNG EMAIL).xlsx'

def clean_email(em):
    if not em:
        return ''
    # remove internal spaces like '. no' -> '.no'
    em = re.sub(r'\s+', '', em.strip().lower())
    return em

def score_row(r):
    score = 0
    status = r.get('status_brreg', '').strip().lower()
    if status == 'aktiv':
        score += 100
    elif status == 'avvikling':
        score -= 50
    elif status == 'konkurs':
        score -= 100
        
    if r.get('contact_name', '').strip():
        score += 50
        
    role = r.get('contact_role', '').strip().lower()
    if 'daglig leder' in role:
        score += 20
    elif 'styrets leder' in role:
        score += 10
        
    if r.get('phone', '').strip() or r.get('phone_all', '').strip():
        score += 20
    if r.get('city', '').strip():
        score += 15
    if r.get('website', '').strip():
        score += 15
        
    form = r.get('legal_form', '').strip().upper()
    if form == 'AS':
        score += 10
    elif form == 'NUF':
        score += 5
        
    return score

with open(INPUT_FILE, 'r', encoding='utf-8-sig', errors='ignore') as f:
    reader = csv.DictReader(f)
    fieldnames = list(reader.fieldnames)
    rows = list(reader)

total = len(rows)
email_groups = {}
no_email_count = 0

for r in rows:
    raw_email = r.get('email', '')
    em = clean_email(raw_email)
    if not em or '@' not in em:
        no_email_count += 1
        continue
    if em not in email_groups:
        email_groups[em] = []
    # Save cleaned email
    r['email'] = em
    email_groups[em].append(r)

# Select best row for each unique email
deduped_rows = []
for em, group in email_groups.items():
    group.sort(key=score_row, reverse=True)
    best = group[0]
    # Keep track of duplicate count
    best['branch_count'] = len(group)
    deduped_rows.append(best)

# Sort: Active first, then by whether contact name exists, then by company name
deduped_rows.sort(key=lambda r: (
    0 if r.get('status_brreg', '').strip().lower() == 'aktiv' else 1,
    0 if r.get('contact_name', '').strip() else 1,
    r.get('name', '').strip().lower()
))

# Filter active only for cold mail recommendation
active_only = [r for r in deduped_rows if r.get('status_brreg', '').strip().lower() == 'aktiv']
inactive_only = [r for r in deduped_rows if r.get('status_brreg', '').strip().lower() != 'aktiv']

print(f"Tổng số dòng ban đầu: {total:,}")
print(f"Số dòng không có email: {no_email_count:,}")
print(f"Số email duy nhất sau lọc trùng: {len(deduped_rows):,}")
print(f"  • Doanh nghiệp đang hoạt động (Aktiv): {len(active_only):,}")
print(f"  • Doanh nghiệp phá sản/giải thể: {len(inactive_only):,}")

# Write to CSV
out_fieldnames = list(fieldnames)
if 'branch_count' not in out_fieldnames:
    out_fieldnames.append('branch_count')

with open(OUTPUT_CSV, 'w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=out_fieldnames)
    writer.writeheader()
    writer.writerows(deduped_rows)

print(f"[+] Đã lưu CSV lọc trùng tại: {OUTPUT_CSV}")

# Write to Excel with formatting
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Unique Cold Mail Leads"
ws.views.sheetView[0].showGridLines = True

font_family = "Segoe UI"
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
data_font = Font(name=font_family, size=10, bold=False, color="000000")

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

# Headers to display
ws.append(out_fieldnames)
for col_idx in range(1, len(out_fieldnames) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

for r_idx, r in enumerate(deduped_rows, 2):
    row_data = [r.get(col, '') for col in out_fieldnames]
    ws.append(row_data)
    
    is_active = r.get('status_brreg', '').strip().lower() == 'aktiv'
    row_fill = None if is_active else PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # light red for inactive
    
    for col_idx in range(1, len(out_fieldnames) + 1):
        cell = ws.cell(row=r_idx, column=col_idx)
        cell.font = data_font
        cell.border = thin_border
        if row_fill:
            cell.fill = row_fill
            
        col_name = out_fieldnames[col_idx - 1]
        if col_name in ['orgnr', 'phone', 'phone_all', 'postal_code', 'employees', 'revenue', 'profit', 'status_brreg', 'branch_count']:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Auto-adjust column widths
for col in ws.columns:
    max_len = 0
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    for cell in col[:100]: # sample first 100 rows for speed
        val_str = str(cell.value or '')
        if val_str.startswith("'"):
            val_str = val_str[1:]
        max_len = max(max_len, len(val_str))
    ws.column_dimensions[col_letter].width = max(min(max_len + 3, 40), 12)

wb.save(OUTPUT_XLSX)
print(f"[+] Đã lưu Excel format đẹp tại: {OUTPUT_XLSX}")
