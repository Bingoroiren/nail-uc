# -*- coding: utf-8 -*-
"""
Chuẩn Hóa Danh Sách Agentúra dočasného zamestnávania (ADZ) Từ ÚPSVaR Slovakia
Đầu vào: SK_ADZC_UPSV_0RAW.xlsx
Đầu ra:
  1. SK_ADZC_UPSV_2COL.csv (Chuẩn 20 cột Cold Mail, 100% Email đẩy lên đầu, SĐT có nháy đơn)
  2. data/formatted/SK_ADZC_UPSV_2COL.csv
  3. data/formatted/SK_ADZC_UPSV_1ENR.csv (Bản chi tiết lưu giữ IČO, giấy phép, địa bàn)
"""

import os
import sys
import re
import csv
import shutil
import openpyxl
import urllib.parse

# Đảm bảo UTF-8 cho Windows Console
if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8', line_buffering=True)
        sys.stderr.reconfigure(encoding='utf-8', line_buffering=True)
    except Exception:
        pass

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT_XLSX = os.path.join(ROOT_DIR, "SK_ADZC_UPSV_0RAW.xlsx")
OUTPUT_COLDMAIL_CSV = os.path.join(ROOT_DIR, "SK_AUTO_ZLAT_2COL.csv") # placeholder
ROOT_COLDMAIL_CSV = os.path.join(ROOT_DIR, "SK_ADZC_UPSV_2COL.csv")
FMT_COLDMAIL_CSV = os.path.join(ROOT_DIR, "data", "formatted", "SK_ADZC_UPSV_2COL.csv")
FMT_DETAIL_CSV = os.path.join(ROOT_DIR, "data", "formatted", "SK_ADZC_UPSV_1ENR.csv")

COLDMAIL_FIELDNAMES = [
    "No.",
    "Công ty",
    "Chức danh",
    "Người liên hệ",
    "SĐT",
    "Liên Hệ",
    "Email",
    "Liên Hệ mail",
    "Địa chỉ",
    "Lương",
    "Ngày đăng",
    "Hạn tuyển",
    "Check gửi",
    "Last Subject",
    "Last Body HTML",
    "Trạng thái Reply",
    "Lần Follow-up",
    "Ngày Follow-up gần nhất",
    "Mailbox đã dùng",
    "Category"
]

DETAIL_FIELDNAMES = [
    "No.",
    "Evid_Cislo",
    "Cislo_Povolenia",
    "Nazov_Spolocnosti",
    "ICO",
    "Ulica_Cislo",
    "PSC",
    "Mesto",
    "Adresa_Komplet",
    "Telefon",
    "Mobil",
    "SDT_Chuan",
    "Email",
    "Web",
    "Datum_Vydania",
    "Datum_Pravoplatnosti",
    "Okruh_Zamestnani",
    "Uzemny_Rozsah",
    "Category"
]

def clean_company_name(name):
    if not name:
        return ""
    name = re.sub(r'[\r\n\t]+', ' ', str(name))
    name = re.sub(r'\s+', ' ', name).strip()
    return name

def clean_email(raw_email):
    if not raw_email:
        return ""
    e = urllib.parse.unquote(str(raw_email)).replace('%20', '').strip().lower()
    e = re.sub(r'^(e-mail:|email:|mail:)\s*', '', e)
    # Nếu có nhiều mail ngăn cách bởi dấu phẩy, chấm phẩy, khoảng trắng
    emails = re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', e)
    if not emails:
        return ""
    
    # Ưu tiên email theo chuẩn B2B Agency Manpower
    def score_email(em):
        u, d = em.split('@', 1)
        score = 10
        if any(hr in u for hr in ['praca', 'kariera', 'hr', 'jobs', 'nabor']):
            score += 50
        elif any(mgr in u for mgr in ['vedenie', 'riaditel', 'konatel', 'director', 'manager', 'ceo', 'sales']):
            score += 45
        elif any(info in u for info in ['info', 'kontakt', 'contact', 'office', 'biuro']):
            score += 40
        elif d in ['gmail.com', 'seznam.cz', 'zoznam.sk']:
            score += 15
        return score

    best_em = max(emails, key=score_email)
    return best_em.rstrip('.,;:')

def extract_contact_info(raw_tel, raw_mob, raw_email):
    # 1. Bóc tách Email
    all_text = f"{raw_tel or ''} {raw_mob or ''} {raw_email or ''}"
    emails = re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', all_text)
    
    clean_email = ""
    if emails:
        # B2B scoring
        def score_email(em):
            u, d = em.lower().split('@', 1)
            score = 10
            if any(hr in u for hr in ['praca', 'kariera', 'hr', 'jobs', 'nabor']):
                score += 50
            elif any(mgr in u for mgr in ['vedenie', 'riaditel', 'konatel', 'director', 'manager', 'ceo', 'sales']):
                score += 45
            elif any(info in u for info in ['info', 'kontakt', 'contact', 'office', 'biuro']):
                score += 40
            elif d in ['gmail.com', 'seznam.cz', 'zoznam.sk']:
                score += 15
            return score
        clean_email = max(emails, key=score_email).lower().rstrip('.,;:')

    # 2. Bóc tách Số điện thoại
    phone_text = f"{raw_tel or ''} {raw_mob or ''}"
    if clean_email:
        phone_text = phone_text.replace(clean_email, ' ')
    # Bỏ các nhãn text
    phone_text = re.sub(r'(tel/fax:|tel:|mobil:|mob:|fax:)', ' ', phone_text, flags=re.I)
    
    # Tách các số điện thoại theo dấu phẩy, chấm phẩy, xuống dòng
    parts = re.split(r'[,;\n\r]+', phone_text)
    clean_phones = []
    for p in parts:
        p = p.strip()
        digits = re.sub(r'\D', '', p)
        if len(digits) >= 6:
            p_formatted = re.sub(r'\s+', ' ', p)
            clean_phones.append(p_formatted)
            
    final_phone = clean_phones[0] if clean_phones else ""
    if final_phone:
        while final_phone.startswith("'"):
            final_phone = final_phone[1:]
        final_phone = f"'{final_phone}"

    return final_phone, clean_email

def normalize_web(w):
    if not w or str(w).strip().lower() in ['none', '-']:
        return ""
    w = str(w).strip()
    if not w.startswith('http://') and not w.startswith('https://'):
        w = f"https://{w}"
    return w

def process_adzc_file():
    if not os.path.exists(INPUT_XLSX):
        print(f"[ERROR] Không tìm thấy file {INPUT_XLSX}")
        return

    print(f"[*] Đang đọc file Excel ÚPSVaR: {INPUT_XLSX}...")
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    
    if 'AKTÍVNE' not in wb.sheetnames:
        print("[ERROR] Không tìm thấy sheet AKTÍVNE")
        return
        
    ws = wb['AKTÍVNE']
    print(f"[*] Sheet AKTÍVNE có tổng cộng {ws.max_row} dòng.")

    active_records = []
    seen_companies = set()

    for r in range(6, ws.max_row + 1):
        # Chỉ lấy các dòng đang hoạt động (không bị ẩn bởi AutoFilter)
        if ws.row_dimensions[r].hidden:
            continue

        zrus = ws.cell(r, 4).value
        # Nếu có ngày hủy thì bỏ qua
        if zrus and str(zrus).strip():
            continue

        name_legal = ws.cell(r, 6).value
        name_person = ws.cell(r, 12).value
        name = clean_company_name(name_legal or name_person)

        if not name or name.startswith('Názov') or name.startswith('Evid.'):
            continue

        ico = str(ws.cell(r, 7).value or '').strip()
        street = str(ws.cell(r, 8).value or '').strip()
        num = str(ws.cell(r, 9).value or '').strip()
        zip_c = str(ws.cell(r, 10).value or '').strip()
        city = str(ws.cell(r, 11).value or '').strip()

        ulica_cislo = f"{street} {num}".strip()
        addr_parts = [p for p in [ulica_cislo, zip_c, city] if p]
        address = ', '.join(addr_parts)

        tel = ws.cell(r, 21).value
        mob = ws.cell(r, 22).value
        raw_email = ws.cell(r, 24).value
        phone, email = extract_contact_info(tel, mob, raw_email)

        raw_web = ws.cell(r, 25).value
        web = normalize_web(raw_web)

        # Lọc bỏ nếu công ty KHÔNG CÓ THÔNG TIN LIÊN LẠC (Không SĐT và Không Email)
        if not phone and not email:
            continue

        # Khử trùng lặp theo Tên hoặc IČO
        dedup_key = ico if ico else name.lower()
        if dedup_key in seen_companies:
            continue
        seen_companies.add(dedup_key)

        evid_cislo = str(ws.cell(r, 1).value or '').strip()
        cislo_povol = str(ws.cell(r, 2).value or '').strip()
        datum_vyd = str(ws.cell(r, 17).value or '').strip()
        datum_prav = str(ws.cell(r, 18).value or '').strip()
        okruh = str(ws.cell(r, 19).value or '').strip()
        uzemie = str(ws.cell(r, 20).value or '').strip()

        active_records.append({
            'name': name,
            'ico': ico,
            'address': address,
            'phone': phone,
            'email': email,
            'web': web,
            'evid_cislo': evid_cislo,
            'cislo_povol': cislo_povol,
            'datum_vyd': datum_vyd,
            'datum_prav': datum_prav,
            'okruh': okruh,
            'uzemie': uzemie,
            'street_num': ulica_cislo,
            'zip': zip_c,
            'city': city
        })

    print(f"[*] Đã lọc được {len(active_records)} Agency đang hoạt động có đầy đủ SĐT/Email.")

    # Sắp xếp: Đẩy 100% công ty CÓ EMAIL lên đầu
    with_email = [r for r in active_records if r['email']]
    without_email = [r for r in active_records if not r['email']]
    sorted_records = with_email + without_email

    print(f"[*] Phân loại dữ liệu:")
    print(f"    - Có Email (Xếp đầu dòng 1 đến {len(with_email)}): {len(with_email)} công ty")
    print(f"    - Có Số điện thoại (Xếp sau): {len(without_email)} công ty")

    # 1. Xuất file Cold Mail chuẩn 20 cột
    coldmail_rows = []
    category_label = "Môi giới việc làm / Agentúra dočasného zamestnávania (ADZ)"

    for i, r in enumerate(sorted_records, start=1):
        coldmail_rows.append({
            "No.": str(i),
            "Công ty": r['name'],
            "Chức danh": "",
            "Người liên hệ": "",
            "SĐT": r['phone'],
            "Liên Hệ": r['web'] if r['web'] else "",
            "Email": r['email'],
            "Liên Hệ mail": "",
            "Địa chỉ": r['address'],
            "Lương": "",
            "Ngày đăng": "",
            "Hạn tuyển": "",
            "Check gửi": "",
            "Last Subject": "",
            "Last Body HTML": "",
            "Trạng thái Reply": "",
            "Lần Follow-up": "",
            "Ngày Follow-up gần nhất": "",
            "Mailbox đã dùng": "",
            "Category": category_label
        })

    # Ghi file Cold Mail ra root
    with open(ROOT_COLDMAIL_CSV, mode='w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=COLDMAIL_FIELDNAMES)
        writer.writeheader()
        writer.writerows(coldmail_rows)
    print(f"[SUCCESS] Đã xuất file Cold Mail chuẩn 20 cột tại thư mục gốc: {ROOT_COLDMAIL_CSV}")

    # Đồng bộ sang data/formatted/
    os.makedirs(os.path.dirname(FMT_COLDMAIL_CSV), exist_ok=True)
    shutil.copy2(ROOT_COLDMAIL_CSV, FMT_COLDMAIL_CSV)
    print(f"[SUCCESS] Đã đồng bộ sang {FMT_COLDMAIL_CSV}")

    # 2. Xuất file Chi Tiết (Full Metadata kèm IČO, giấy phép) sang data/formatted/
    detail_rows = []
    for i, r in enumerate(sorted_records, start=1):
        detail_rows.append({
            "No.": str(i),
            "Evid_Cislo": r['evid_cislo'],
            "Cislo_Povolenia": r['cislo_povol'],
            "Nazov_Spolocnosti": r['name'],
            "ICO": r['ico'],
            "Ulica_Cislo": r['street_num'],
            "PSC": r['zip'],
            "Mesto": r['city'],
            "Adresa_Komplet": r['address'],
            "Telefon": r['phone'],
            "Mobil": "",
            "SDT_Chuan": r['phone'],
            "Email": r['email'],
            "Web": r['web'],
            "Datum_Vydania": r['datum_vyd'],
            "Datum_Pravoplatnosti": r['datum_prav'],
            "Okruh_Zamestnani": r['okruh'],
            "Uzemny_Rozsah": r['uzemie'],
            "Category": category_label
        })

    with open(FMT_DETAIL_CSV, mode='w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=DETAIL_FIELDNAMES)
        writer.writeheader()
        writer.writerows(detail_rows)
    print(f"[SUCCESS] Đã xuất file chi tiết tra cứu: {FMT_DETAIL_CSV}")

if __name__ == "__main__":
    process_adzc_file()
