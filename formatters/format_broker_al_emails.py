import csv
import os
import shutil
import re
import sys
import urllib.parse
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# Set console output encoding to UTF-8
if sys.platform.startswith('win') and hasattr(sys.stdout, 'buffer'):
    import codecs
    try:
        sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
        sys.stderr = codecs.getwriter('utf-8')(sys.stderr.buffer, 'strict')
    except Exception:
        pass

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(SCRIPT_DIR)

CANDIDATES = [
    os.path.join(ROOT_DIR, "data", "formatted", "broker_albania_with_emails.csv"),
    os.path.join(ROOT_DIR, "data", "raw", "broker_albania.csv"),
    os.path.join(ROOT_DIR, "broker_albania_with_emails.csv"),
    os.path.join(ROOT_DIR, "broker_albania.csv")
]

INPUT_CSV = None
for candidate in CANDIDATES:
    if os.path.exists(candidate):
        INPUT_CSV = candidate
        break

if len(sys.argv) > 1:
    INPUT_CSV = sys.argv[1]
elif not INPUT_CSV:
    INPUT_CSV = CANDIDATES[0]

if INPUT_CSV:
    ext = os.path.splitext(INPUT_CSV)[1]
    base_name = os.path.basename(INPUT_CSV)
    BACKUP_CSV = os.path.join(ROOT_DIR, "data", "formatted", base_name.replace(ext, f"_backup{ext}"))
    FORMATTED_CSV = os.path.join(ROOT_DIR, "data", "formatted", "broker_albania_final.csv")
    FORMATTED_XLSX = os.path.join(ROOT_DIR, "data", "formatted", "broker_albania_final.xlsx")
    FORMATTED_V2_CSV = os.path.join(ROOT_DIR, "data", "formatted", "broker_albania_final_v2.csv")
    FORMATTED_V2_XLSX = os.path.join(ROOT_DIR, "data", "formatted", "broker_albania_final_v2.xlsx")
    ROOT_COPY_CSV = os.path.join(ROOT_DIR, "broker_albania_final.csv")
    ROOT_COPY_XLSX = os.path.join(ROOT_DIR, "broker_albania_final.xlsx")
else:
    BACKUP_CSV = None
    FORMATTED_CSV = None
    FORMATTED_XLSX = None
    FORMATTED_V2_CSV = None
    FORMATTED_V2_XLSX = None
    ROOT_COPY_CSV = None
    ROOT_COPY_XLSX = None

# Category translations from Albanian to Vietnamese
CATEGORY_TRANSLATIONS = {
    "shërbim konsulent për burime njerëzore": "Dịch vụ tư vấn nhân sự",
    "sherbim konsulent per burime njerezore": "Dịch vụ tư vấn nhân sự",
    "agjenci punësimi": "Công ty môi giới việc làm",
    "agjenci punesimi": "Công ty môi giới việc làm",
    "qendra e punësimit": "Trung tâm dịch vụ việc làm",
    "qendra e punesimit": "Trung tâm dịch vụ việc làm",
    "rekrutues": "Nhà tuyển dụng / Săn đầu người",
    "agjenci për punë të përkohshme": "Công ty cung ứng lao động tạm thời",
    "agjenci per pune te perkohshme": "Công ty cung ứng lao động tạm thời"
}

GENERIC_DOMAINS = {
    'gmail.com', 'yahoo.com', 'hotmail.com', 'outlook.com', 'live.com', 
    'aol.com', 'icloud.com', 'mail.com', 'ymail.com', 'msn.com', 
    'wix.com', 'squarespace.com', 'wordpress.com', 'wixpress.com'
}

BUSINESS_PREFIXES = {
    'info', 'contact', 'hello', 'office', 'admin', 'sales', 
    'recruitment', 'hr', 'punesim', 'rekrutim', 'support'
}

BAD_KEYWORDS = {
    'wix', 'support@wix', 'no-reply', 'noreply', 'test@', 'example@', 'domain.com', 'sentry.io'
}

SOCIAL_DOMAINS = {
    'facebook.com', 'fb.com', 'fb.me', 'instagram.com', 'instagr.am',
    'twitter.com', 'x.com', 'linkedin.com', 'youtube.com', 'youtu.be',
    'tiktok.com', 'pinterest.com', 'pin.it', 'reddit.com', 'threads.net',
    't.me', 'telegram.org', 'wa.me', 'whatsapp.com', 'viber.com',
    'google.com', 'goo.gl', 'maps.app.goo.gl'
}

def normalize_text(text):
    if not text:
        return ""
    text = text.lower().strip()
    replacements = {'ë': 'e', 'ç': 'c', 'é': 'e', 'è': 'e'}
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
    return text

def translate_category(raw_cat):
    if not raw_cat:
        return "Môi giới việc làm Albania"
    raw_lower = raw_cat.strip().lower()
    raw_norm = normalize_text(raw_cat)
    
    for tag, vi_trans in CATEGORY_TRANSLATIONS.items():
        tag_norm = normalize_text(tag)
        if tag in raw_lower or tag_norm in raw_norm:
            return vi_trans
            
    return raw_cat.strip()

def get_field(r, *keys):
    for k in keys:
        val = r.get(k, '')
        if val:
            return val
    return ''

def normalize_name(name):
    if not name:
        return ""
    cleaned = re.sub(r'[\(\[\{].*?[\)\]\}]', '', name)
    cleaned = re.sub(r'\b(sh\.p\.k|shpk|ltd|llc|inc|corp|gmbh|sa|d\.o\.o)\b', '', cleaned, flags=re.I)
    cleaned = re.sub(r'[^\w\s]', '', cleaned)
    return " ".join(cleaned.lower().split())

def guess_email_from_website(url):
    if not url:
        return ""
    if not url.startswith(('http://', 'https://')):
        url = 'http://' + url
    try:
        parsed = urllib.parse.urlparse(url)
        netloc = parsed.netloc.strip().lower()
        if not netloc:
            return ""
        if ':' in netloc:
            netloc = netloc.split(':')[0]
        netloc = re.sub(r'^www\d*\.', '', netloc)
        if not netloc or '.' not in netloc:
            return ""
        for social in SOCIAL_DOMAINS:
            if netloc == social or netloc.endswith('.' + social):
                return ""
        if re.match(r'^[a-z0-9][a-z0-9\.\-]*\.[a-z]{2,}$', netloc):
            if netloc.endswith('.') or re.match(r'^\d+\.\d+\.\d+\.\d+$', netloc):
                return ""
            return f"info@{netloc}"
        return ""
    except Exception:
        return ""

def clean_and_score_email(email_str):
    if not email_str:
        return ""
    
    discard_domains = {
        'example.com', 'example.org', 'example.net', 'yourdomain.com', 
        'email.com', 'domain.com', 'website.com', 'company.com', 
        'sentry.io', 'git.com', 'github.com', 'test.com', 'g.co'
    }
    
    emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,7}\b', email_str)
    if not emails:
        return ""
        
    scored = []
    for em in emails:
        em_lower = em.strip().lower()
        
        if any(kw in em_lower for kw in BAD_KEYWORDS):
            continue
            
        parts = em_lower.split('@')
        if len(parts) != 2:
            continue
            
        username, domain = parts
        if domain in discard_domains or domain.endswith('.png') or domain.endswith('.jpg'):
            continue
            
        score = 0
        if any(username.startswith(prefix) for prefix in BUSINESS_PREFIXES):
            score += 10
        if domain not in GENERIC_DOMAINS:
            score += 5
        scored.append((score, em))
        
    if not scored:
        return ""
        
    scored.sort(key=lambda x: x[0], reverse=True)
    # Return highest scoring unique emails separated by comma
    seen = set()
    best_emails = []
    for score, em in scored:
        if em.lower() not in seen:
            seen.add(em.lower())
            best_emails.append(em)
            
    return ", ".join(best_emails)

def format_phone(phone_str):
    if not phone_str:
        return ""
    p = phone_str.strip()
    while p.startswith("'"):
        p = p[1:]
    if not p:
        return ""
    return f"'{p}"

def main():
    if not INPUT_CSV or not os.path.exists(INPUT_CSV):
        print(f"[-] Input file {INPUT_CSV} does not exist.")
        return
        
    os.makedirs(os.path.dirname(FORMATTED_CSV), exist_ok=True)
    if BACKUP_CSV:
        print(f"[*] Backing up original CSV to {BACKUP_CSV}...")
        try:
            shutil.copy2(INPUT_CSV, BACKUP_CSV)
        except Exception:
            pass
    
    rows = []
    with open(INPUT_CSV, mode='r', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for r in reader:
            rows.append(r)
            
    print(f"[+] Loaded {len(rows)} rows from {INPUT_CSV}.")
    
    # Filter closed businesses
    active_rows = []
    closed_count = 0
    for r in rows:
        if r.get('Permanently_Closed', '').strip().lower() == 'yes':
            closed_count += 1
            continue
        active_rows.append(r)
    rows = active_rows
    if closed_count > 0:
        print(f"[*] Filtered out {closed_count} permanently closed businesses.")

    # Step 1: Normalize emails
    for r in rows:
        original_email = get_field(r, 'Email', 'Emails', 'Liên Hệ mail', 'email', 'Email liên hệ')
        cleaned_em = clean_and_score_email(original_email)
        if not cleaned_em:
            website = get_field(r, 'Website', 'website', 'URL', 'Web')
            cleaned_em = guess_email_from_website(website)
        r['Best_Email'] = cleaned_em
        
    # Step 2: Deduplicate by normalized name
    grouped_rows = {}
    for r in rows:
        norm_name = normalize_name(get_field(r, 'Name', 'business_name', 'name', 'Tên công ty'))
        if not norm_name:
            continue
        if norm_name not in grouped_rows:
            grouped_rows[norm_name] = []
        grouped_rows[norm_name].append(r)
        
    deduplicated_rows = []
    for norm_name, r_list in grouped_rows.items():
        def sort_rep(r):
            has_email = 1 if r['Best_Email'] else 0
            has_web = 1 if get_field(r, 'Website', 'website', 'Web').strip() else 0
            try:
                reviews = float(r.get('Reviews_Count') or 0)
            except ValueError:
                reviews = 0.0
            try:
                rating = float(r.get('Rating') or 0)
            except ValueError:
                rating = 0.0
            return (-has_email, -has_web, -reviews, -rating)
            
        r_list.sort(key=sort_rep)
        deduplicated_rows.append(r_list[0])
        
    print(f"[*] Deduplicated by company name: {len(rows)} down to {len(deduplicated_rows)} unique companies.")
    
    # Step 3: Format output records and translate tags
    output_rows = []
    for idx, r in enumerate(deduplicated_rows, 1):
        raw_cat = get_field(r, 'Category', 'category', 'Ngành nghề')
        translated_cat = translate_category(raw_cat)
        
        phone = format_phone(get_field(r, 'Phone', 'phone', 'SĐT'))
        website = get_field(r, 'Website', 'website', 'Web').strip()
        address = get_field(r, 'Address', 'address', 'Địa chỉ').strip()
        state = get_field(r, 'State', 'Location_Name', 'state', 'Khu vực').strip()
        rating = get_field(r, 'Rating', 'rating')
        reviews = get_field(r, 'Reviews_Count', 'reviews_count')
        map_url = get_field(r, 'URL', 'url', 'Maps_URL')
        
        output_rows.append({
            "No.": idx,
            "Tên công ty": get_field(r, 'Name', 'business_name', 'name').strip(),
            "Email": r['Best_Email'],
            "SĐT": phone,
            "Ngành nghề (Dịch)": translated_cat,
            "Tag gốc (Albania)": raw_cat,
            "Địa chỉ": address,
            "Thành phố / Hạt": state,
            "Website": website,
            "Đánh giá": rating,
            "Số lượt đánh giá": reviews,
            "URL Maps": map_url
        })
        
    # Sort: records with email first
    output_rows.sort(key=lambda x: (0 if x["Email"] else 1, x["Tên công ty"].lower()))
    for idx, r in enumerate(output_rows, 1):
        r["No."] = idx
        
    fieldnames = [
        "No.", "Tên công ty", "Email", "SĐT", "Ngành nghề (Dịch)", 
        "Tag gốc (Albania)", "Địa chỉ", "Thành phố / Hạt", "Website", 
        "Đánh giá", "Số lượt đánh giá", "URL Maps"
    ]
    
    # Write CSV
    target_write_file = FORMATTED_CSV
    try:
        with open(FORMATTED_CSV, mode='w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(output_rows)
    except PermissionError:
        target_write_file = FORMATTED_V2_CSV
        print(f"[!] {FORMATTED_CSV} is locked. Writing to {FORMATTED_V2_CSV} instead...")
        with open(FORMATTED_V2_CSV, mode='w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(output_rows)
            
    # Write Excel
    target_xlsx = FORMATTED_XLSX
    try:
        write_excel(output_rows, fieldnames, FORMATTED_XLSX)
    except PermissionError:
        target_xlsx = FORMATTED_V2_XLSX
        print(f"[!] {FORMATTED_XLSX} is locked. Writing to {FORMATTED_V2_XLSX} instead...")
        write_excel(output_rows, fieldnames, FORMATTED_V2_XLSX)
        
    # Also write a copy to project root for convenience
    try:
        if ROOT_COPY_CSV:
            shutil.copy2(target_write_file, ROOT_COPY_CSV)
        if ROOT_COPY_XLSX:
            shutil.copy2(target_xlsx, ROOT_COPY_XLSX)
    except Exception:
        pass
        
    with_email_count = sum(1 for r in output_rows if r["Email"])
    without_email_count = len(output_rows) - with_email_count
    
    print(f"\n================ SUMMARY ================")
    print(f"Tổng số công ty:       {len(output_rows)}")
    print(f"Có Email liên hệ:      {with_email_count}")
    print(f"Chưa có Email:         {without_email_count}")
    print(f"File CSV hoàn chỉnh:   {target_write_file}")
    print(f"File Excel hoàn chỉnh: {target_xlsx}")
    print(f"=========================================\n")

def write_excel(rows, fieldnames, file_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Albania Labor Brokers"
    
    ws.views.sheetView[0].showGridLines = True
    
    font_family = "Segoe UI"
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=11, bold=False, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    # Write header
    ws.append(fieldnames)
    for col_idx in range(1, len(fieldnames) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border
        
    # Write rows
    for r_idx, r in enumerate(rows, 2):
        row_data = [r[col] for col in fieldnames]
        ws.append(row_data)
        
        has_email = bool(r["Email"])
        row_fill = PatternFill(start_color="EBF1F5", end_color="EBF1F5", fill_type="solid") if has_email else None
        
        for col_idx in range(1, len(fieldnames) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
                
            col_name = fieldnames[col_idx - 1]
            if col_name in ["No.", "SĐT", "Đánh giá", "Số lượt đánh giá"]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
            if col_name == "SĐT" and cell.value:
                cell.number_format = '@'
                
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if val_str.startswith("'"):
                val_str = val_str[1:]
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)
        
    wb.save(file_path)

if __name__ == "__main__":
    main()
